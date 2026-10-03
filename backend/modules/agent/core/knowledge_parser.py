#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
知识库文档解析器

按扩展名分发到具体解析实现，统一返回纯文本。新增格式 = 在 _PARSERS 注册新解析函数。
pdf/docx 为可选依赖，缺包时惰性导入报错提示补装。
"""
import logging
from io import BytesIO
from typing import Callable

from core.exception.errors import CustomError
from core.i18n import t
from core.response.response_code import CustomErrorCode

logger = logging.getLogger(__name__)

# 支持的扩展名（小写、无点）
SUPPORTED_TYPES = {"txt", "md", "pdf", "docx"}


def _parse_plain(content: bytes) -> str:
    """txt / md：UTF-8 宽松解码"""
    return content.decode("utf-8", errors="replace")


def _parse_pdf(content: bytes) -> str:
    """pdf：pypdf 逐页提取"""
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise CustomError(
            error=CustomErrorCode.KNOWLEDGE_FILE_TYPE_UNSUPPORTED,
            msg="pypdf 未安装（uv add pypdf 后重启）",
        ) from exc
    reader = PdfReader(BytesIO(content))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n\n".join(pages)


def _parse_docx(content: bytes) -> str:
    """docx：python-docx 逐段提取（含表格单元格）"""
    try:
        import docx
    except ImportError as exc:
        raise CustomError(
            error=CustomErrorCode.KNOWLEDGE_FILE_TYPE_UNSUPPORTED,
            msg="python-docx 未安装（uv add python-docx 后重启）",
        ) from exc
    document = docx.Document(BytesIO(content))
    parts = [para.text for para in document.paragraphs if para.text]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


_PARSERS: dict[str, Callable[[bytes], str]] = {
    "txt": _parse_plain,
    "md": _parse_plain,
    "pdf": _parse_pdf,
    "docx": _parse_docx,
}


def extract_file_type(file_name: str) -> str:
    """从文件名取小写扩展名（无点）；无扩展名返回空串"""
    name = (file_name or "").rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
    if "." not in name:
        return ""
    return name.rsplit(".", 1)[-1].lower()


def parse_document(file_name: str, content: bytes) -> str:
    """
    解析文档为纯文本。

    - 扩展名白名单外抛 KNOWLEDGE_FILE_TYPE_UNSUPPORTED
    - 提取文本清洗后为空抛 KNOWLEDGE_EMPTY_CONTENT
    """
    file_type = extract_file_type(file_name)
    parser = _PARSERS.get(file_type)
    if parser is None:
        raise CustomError(
            error=CustomErrorCode.KNOWLEDGE_FILE_TYPE_UNSUPPORTED,
            msg=t("error.knowledge.file_type_unsupported", name=file_name),
        )
    text = parser(content)
    # 压缩连续空白行（保留段落结构），去掉每行首尾空白
    lines = [line.strip() for line in text.splitlines()]
    cleaned: list[str] = []
    blank = 0
    for line in lines:
        if line:
            cleaned.append(line)
            blank = 0
        else:
            blank += 1
            if blank <= 1:
                cleaned.append("")
    text = "\n".join(cleaned).strip()
    if not text:
        raise CustomError(
            error=CustomErrorCode.KNOWLEDGE_EMPTY_CONTENT,
            msg=t("error.knowledge.empty_content", name=file_name),
        )
    return text
